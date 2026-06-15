import os
import zlib
import io
from PIL import Image
from pypdf import PdfReader, PdfWriter
from docx import Document
from openpyxl import load_workbook


def compress_image(file_bytes: bytes, ext: str) -> bytes:
    extension = ext.lower()
    try:
        img_io = io.BytesIO(file_bytes)
        img = Image.open(img_io)
        output_io = io.BytesIO()

        if extension in [".jpg", ".jpeg"]:
            if img.mode in ("RGBA", "LA"):
                img = img.convert("RGB")
            img.save(output_io, format="JPEG", quality=60)
            return output_io.getvalue()

        elif extension == ".png":
            img.save(output_io, format="PNG", optimize=True, compress_level=7)
            return output_io.getvalue()

        elif extension == ".webp":
            img.save(output_io, format="WEBP", quality=60)
            return output_io.getvalue()

        return file_bytes
    except Exception as e:
        print(f"Gagal mengompresi gambar: {e}")
        return file_bytes


def compress_pdf(file_bytes: bytes) -> bytes:
    try:
        input_stream = io.BytesIO(file_bytes)
        reader = PdfReader(input_stream)
        writer = PdfWriter()

        for page in reader.pages:
            page.compress_content_streams()
            writer.add_page(page)

        output_stream = io.BytesIO()
        writer.write(output_stream)
        return output_stream.getvalue()
    except Exception as e:
        print(f"Gagal mengompresi PDF: {e}")
        return file_bytes


def compress_docx(file_bytes: bytes) -> bytes:
    try:
        doc_stream = io.BytesIO(file_bytes)
        doc = Document(doc_stream)

        for rel in doc.part.relations.values():
            if "image" in rel.target_ref:
                image_part = rel.target_part

                img = Image.open(io.BytesIO(image_part.blob))
                compressed_img_io = io.BytesIO()

                img_format = img.format if img.format else "JPEG"
                if img.mode in ("RGBA", "LA") and img_format == "JPEG":
                    img = img.convert("RGB")

                img.save(compressed_img_io, format=img_format, quality=60)
                image_part._blob = compressed_img_io.getvalue()

        output_stream = io.BytesIO()
        doc.save(output_stream)
        return output_stream.getvalue()
    except Exception as e:
        print(f"Gagal mengompresi Word: {e}")
        return file_bytes


def compress_xlsx(file_bytes: bytes) -> bytes:
    try:
        wb_stream = io.BytesIO(file_bytes)
        wb = load_workbook(wb_stream)

        output_stream = io.BytesIO()
        wb.save(output_stream)
        return output_stream.getvalue()
    except Exception as e:
        print(f"Gagal mengompresi Excel: {e}")
        return file_bytes


def compress_file(file_bytes: bytes, file_name: str) -> dict:
    _, ext = os.path.splitext(file_name)
    ext = ext.lower()

    if ext in [".jpg", ".jpeg", ".png", ".webp", ".tiff", ".bmp"]:
        return {
            "buffer": compress_image(file_bytes, ext),
            "is_compressed": True,
            "encoding": None,
        }
    elif ext == ".pdf":
        return {
            "buffer": compress_pdf(file_bytes),
            "is_compressed": True,
            "encoding": None,
        }

    elif ext in [".docx"]:
        return {
            "buffer": compress_docx(file_bytes),
            "is_compressed": True,
            "encoding": None,
        }

    elif ext in [".xlsx"]:
        return {
            "buffer": compress_xlsx(file_bytes),
            "is_compressed": True,
            "encoding": None,
        }

    try:
        compressor = zlib.compressobj(level=6, wbits=zlib.MAX_WBITS | 16)
        compressed_data = compressor.compress(file_bytes) + compressor.flush()
        return {"buffer": compressed_data, "is_compressed": True, "encoding": "gzip"}
    except Exception as e:
        print(f"Gagal melakukan Gzip pada file umum: {e}")
        return {"buffer": file_bytes, "is_compressed": False, "encoding": None}
